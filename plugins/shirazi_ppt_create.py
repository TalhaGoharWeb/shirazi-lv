"""PowerPoint creation (Shirazi LV plugin).

Gap filled from the MYRA ULTRA survey: Shirazi's file_processor could read
and summarise .pptx files but never create one. This builds a real deck with
python-pptx (already a dependency): a title slide, one content slide per
section with its bullets, and a closing slide. Saved to the Desktop and
opened, the same pattern as shirazi_excel_create.
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

PLUGIN = {
    "name": "shirazi_ppt_create",
    "description": (
        "Builds a REAL PowerPoint .pptx file and saves it to the Desktop. Use when the "
        "user wants an actual presentation FILE — 'presentation bana do', 'slides banao', "
        "'make a deck on photosynthesis'. YOU compose the structure: a title plus 3-8 "
        "slides, each with a heading and 2-5 short bullets. Invent sensible content when "
        "the user only gave a topic. Say one short sentence first so they don't hear "
        "silence while you compose. Never use this just to answer a question verbally."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "title": {"type": "STRING", "description": "Presentation title, in the user's language."},
            "subtitle": {"type": "STRING", "description": "Optional subtitle for the title slide."},
            "slides": {
                "type": "ARRAY",
                "description": "Content slides, in order.",
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "heading": {"type": "STRING", "description": "Slide heading."},
                        "bullets": {
                            "type": "ARRAY",
                            "items": {"type": "STRING"},
                            "description": "Short bullet points for the slide body.",
                        },
                    },
                },
            },
            "filename": {"type": "STRING", "description": "Optional file name without extension."},
        },
        "required": ["title", "slides"],
    },
}

PLUGIN_SETTINGS = None


def _log(player, msg: str) -> None:
    if player:
        try:
            player.write_log(f"SHIRAZI: {msg}")
        except Exception:
            pass


def _slug(text: str, fallback: str = "presentation") -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", text or "").strip("-").lower()
    return (s[:40] or fallback)


def _open_file(path: Path) -> None:
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(path))  # noqa: S606
        elif sys.platform == "darwin":
            import subprocess

            subprocess.Popen(["open", str(path)])
        else:
            import subprocess

            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass


def _layout(prs, preferred: int, fallback: int = 6):
    layouts = prs.slide_layouts
    for idx in (preferred, fallback, len(layouts) - 1):
        try:
            return layouts[max(0, min(idx, len(layouts) - 1))]
        except Exception:
            continue
    return layouts[0]


def _set_title(slide, text: str) -> None:
    try:
        if slide.shapes.title is not None:
            slide.shapes.title.text = text
            return
    except Exception:
        pass
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.idx == 0:
                shape.text = text
                return
        except Exception:
            continue


def _add_content_slide(prs, heading: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(_layout(prs, 1))
    _set_title(slide, heading)
    body = None
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.idx == 1 and shape.has_text_frame:
                body = shape
                break
        except Exception:
            continue
    if body is None:
        return
    tf = body.text_frame
    tf.word_wrap = True
    lines = [b for b in bullets if str(b).strip()]
    if not lines:
        tf.text = ""
        return
    tf.text = str(lines[0])
    for line in lines[1:]:
        tf.add_paragraph().text = str(line)


def _build(title: str, subtitle: str, slides_in: list, filename: str,
           player) -> str:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except Exception:
        return "python-pptx is not installed on this PC, so I cannot build presentations."

    prs = Presentation()
    try:
        prs.slide_width = Inches(13.333)   # 16:9
        prs.slide_height = Inches(7.5)
    except Exception:
        pass

    cover = prs.slides.add_slide(_layout(prs, 0))
    _set_title(cover, title)
    if subtitle:
        for shape in cover.placeholders:
            try:
                if shape.placeholder_format.idx == 1 and shape.has_text_frame:
                    shape.text_frame.text = subtitle
                    break
            except Exception:
                continue

    made = 0
    for s in slides_in:
        if not isinstance(s, dict):
            continue
        heading = str(s.get("heading") or "").strip()
        bullets = s.get("bullets") or []
        if not heading and not bullets:
            continue
        _add_content_slide(prs, heading or title, [str(b) for b in bullets])
        made += 1

    closing = prs.slides.add_slide(_layout(prs, 5, 6))
    _set_title(closing, "Shukriya" if any(ord(c) > 127 for c in title) else "Thank You")

    name = _slug(filename or title)
    desktop = Path.home() / "Desktop"
    out_dir = desktop if desktop.exists() else Path.home()
    out = out_dir / f"{name}.pptx"
    n = 2
    while out.exists():
        out = out_dir / f"{name}-{n}.pptx"
        n += 1
    prs.save(str(out))
    _open_file(out)
    _log(player, f"Presentation saved: {out}")
    return (f"Presentation '{title}' is ready on the Desktop as {out.name} — "
            f"1 title slide and {made} content slide{'s' if made != 1 else ''}.")


def run(parameters: dict, player=None, session_memory=None) -> str:
    """Mark-LV plugin contract: never raise; return a spoken string."""
    try:
        params = dict(parameters or {})
        title = str(params.get("title") or "").strip()
        slides = params.get("slides") or []
        if not title:
            return "What should the presentation be about? Give me a title or a topic."
        if not isinstance(slides, list) or not slides:
            return (f"'{title}' ke liye slides ka outline chahiye — har slide ka "
                    "heading aur uske points. Topic bata dein, main bana dunga.")
        return _build(title, str(params.get("subtitle") or "").strip(),
                      slides, str(params.get("filename") or "").strip(), player)
    except Exception as e:
        return f"Presentation could not be built: {e}"
